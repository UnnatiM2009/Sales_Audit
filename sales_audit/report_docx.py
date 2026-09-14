"""The Word audit document: scorecard, findings, corrective actions, sign-off."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from .config import INPUT_FILES, MISSING_DATA_NOTES, Config
from .loaders import AuditData
from .physical import SECTIONS
from .scoring import AuditResult

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x5C, 0x8A)
FILL_NAVY, FILL_LIGHT, FILL_BAND, FILL_GREY = "1F3864", "DCE6F1", "F2F5FA", "E8E8E8"
FILL_RED, FILL_AMBER, FILL_GREEN = "FFC7CE", "FFEB9C", "C6EFCE"
FONT = "Calibri"


# --- low-level helpers -------------------------------------------------------
def _shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _set_width(cell, twips: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(twips))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def _write(cell, text, *, bold=False, size=9, color=None, align=None, italic=False):
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.space_before = Pt(1)
    if align:
        para.alignment = align
    run = para.add_run(str(text) if text is not None else "")
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return para


def _add_note(cell, text, size=7):
    para = cell.add_paragraph()
    para.paragraph_format.space_after = Pt(1)
    run = para.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.italic = True
    run.font.color.rgb = BLUE


def heading(doc, text, size=13):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = True
    run.font.color.rgb = NAVY
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), FILL_NAVY)
    bottom.set(qn("w:space"), "2")
    borders.append(bottom)
    pPr.append(borders)
    return p


def para(doc, runs, size=9, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if isinstance(runs, str):
        runs = [(runs, False)]
    for text, bold in runs:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.bold = bold
    return p


def bullet(doc, text, size=9):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    return p


def _set_grid(table, widths):
    """Pin the column grid; without this Word and LibreOffice redistribute widths."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(w)))
        grid.append(gc)
    tbl.insert(1, grid)


def make_table(doc, widths, headers, header_size=9):
    t = doc.add_table(rows=1, cols=len(widths))
    t.style = "Table Grid"
    t.autofit = False
    _set_grid(t, widths)
    for j, (w, h) in enumerate(zip(widths, headers)):
        cell = t.rows[0].cells[j]
        _set_width(cell, w)
        _shade(cell, FILL_NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _write(cell, h, bold=True, size=header_size,
               color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
    return t


def rag_fill(rag: str) -> str | None:
    return {"Red": FILL_RED, "Amber": FILL_AMBER, "Green": FILL_GREEN}.get(rag.split(" ")[0])


# --- document ----------------------------------------------------------------
def write_report(path: str | Path, result: AuditResult, data: AuditData, cfg: Config) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Twips(16838), Twips(11906)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Twips(600))
    total_w = 15638

    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(9)

    ctx = result.context
    band = result.band(cfg.grading)

    # --- title ---------------------------------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SALES PROCESS AUDIT")
    r.font.name, r.font.size, r.bold, r.font.color.rgb = FONT, Pt(19), True, NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Weighted process scorecard — enquiry, conversion, volume, manpower, "
                  "stock, profitability, customer experience and physical showroom audit")
    r.font.name, r.font.size, r.italic, r.font.color.rgb = FONT, Pt(9), True, BLUE

    para(doc, [("Purpose. ", True),
               ("This audit scores how well the sales process is performing against benchmarks, "
                "weights each area by business impact, and produces a single Sales Process Index "
                "that can be tracked month on month and compared across branches.", False)])
    para(doc, [("Method. ", True),
               (f"Each parameter is scored as Weight × Achievement %, capped at 100%. "
                f"Achievement is measured against the norms in {cfg.path.name}. Lines whose "
                f"source data was not supplied are marked 'not scored' with the missing file "
                f"named, and are excluded from the index denominator rather than scored zero. "
                f"Any parameter below {cfg.rag['amber_min']}% achievement carries a line in the "
                f"corrective action plan.", False)], space_after=8)

    # --- identification ------------------------------------------------------
    heading(doc, "Audit identification")
    t = make_table(doc, [2600, 5200, 2600, 5238],
                   ["Field", "Detail", "Field", "Detail"])
    t._tbl.remove(t.rows[0]._tr)
    ident = [
        ("Dealership", cfg.dealership.get("name", ""), "Dealer code", cfg.dealership.get("dealer_code", "")),
        ("Branches in scope", f"{len(cfg.branches)} branches", "Audit date", ctx.get("audit_date", "")),
        ("Auditor", "", "Audit period", ctx.get("period", "")),
        ("Population", ctx.get("population", ""), "Sales Process Index",
         f"{ctx.get('index_text', '—')}  —  {band['band']}"),
    ]
    for a, b, c, d in ident:
        row = t.add_row()
        for j, (val, w, is_label) in enumerate([(a, 2600, True), (b, 5200, False),
                                                (c, 2600, True), (d, 5238, False)]):
            cell = row.cells[j]
            _set_width(cell, w)
            if is_label:
                _shade(cell, FILL_LIGHT)
            _write(cell, val, bold=is_label)

    # --- source register -----------------------------------------------------
    heading(doc, "Input source files")
    para(doc, "Every figure in this audit traces back to a file listed here. Files marked "
              "'Not supplied' are the reason some lines remain unscored.", size=8)
    widths = [900, 5200, 1600, 1600, 3000, 3338]
    t = make_table(doc, widths, ["Ref", "Input source file name", "Sheet", "Records",
                                 "What it is", "Status / what it scores"])
    idx = 0
    for spec in INPUT_FILES:
        ds = data.ds(spec.ref)
        vals = ([spec.ref, ds.filename, ds.sheet, f"{ds.rows:,}", spec.label, spec.scores]
                if ds else [spec.ref, spec.pattern, "—", "—", spec.label, "NOT SUPPLIED"])
        row = t.add_row()
        for j, v in enumerate(vals):
            cell = row.cells[j]
            _set_width(cell, widths[j])
            if idx % 2:
                _shade(cell, FILL_BAND)
            _write(cell, v, bold=(j == 0), size=8)
        idx += 1
    for key, label in MISSING_DATA_NOTES.items():
        if cfg.target(key) or cfg.norm(key):
            continue
        row = t.add_row()
        for j, v in enumerate(["—", label, "Not supplied", "—", "Reference data",
                               "Lines left unscored without it"]):
            cell = row.cells[j]
            _set_width(cell, widths[j])
            if idx % 2:
                _shade(cell, FILL_BAND)
            _write(cell, v, bold=(j == 0), size=8)
        idx += 1

    # --- Section A: scorecard -------------------------------------------------
    doc.add_page_break()
    heading(doc, "Section A — Process parameter scorecard")
    para(doc, f"RAG: Green ≥{cfg.rag['green_min']}% achievement, "
              f"Amber {cfg.rag['amber_min']}–{cfg.rag['green_min'] - 1}%, "
              f"Red below {cfg.rag['amber_min']}%.", size=8)
    cw = [520, 1950, 2850, 2250, 3150, 480, 560, 560, 3318]
    t = make_table(doc, cw, ["S.No", "Process parameter", "What to verify / data source",
                             "Benchmark / norm", "Actual observed", "Wt", "Ach", "Score", "RAG / observation"])
    for p_ in result.pillars:
        row = t.add_row()
        merged = row.cells[0]
        for c in row.cells[1:]:
            merged = merged.merge(c)
        _shade(merged, FILL_LIGHT)
        _write(merged, f"{p_.title}  —  weightage {p_.weight:g}  "
                       f"(scored {p_.scorable_weight:.1f}, achievement "
                       f"{p_.achievement:.1f}%)" if p_.achievement is not None
               else f"{p_.title}  —  weightage {p_.weight:g}  (not scored)",
               bold=True, color=NAVY)
        for i, l in enumerate(p_.lines):
            row = t.add_row()
            shade = FILL_BAND if i % 2 else None
            vals = [l.code, l.name, l.checkpoint, l.norm_text, l.actual_text,
                    f"{l.weight:.1f}", l.display_achievement(), l.display_score(),
                    l.rag_remark(cfg.rag)]
            for j, v in enumerate(vals):
                cell = row.cells[j]
                _set_width(cell, cw[j])
                if shade:
                    _shade(cell, shade)
                _write(cell, v, bold=(j in (0, 7)), size=8,
                       align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 5, 6, 7) else None)
                if j == 4 and l.source and l.source != "Not supplied":
                    _add_note(cell, f"Source: {l.source}")
            fill = rag_fill(l.rag(cfg.rag))
            if fill:
                _shade(row.cells[8], fill)

    row = t.add_row()
    merged = row.cells[0]
    for c in row.cells[1:5]:
        merged = merged.merge(c)
    _shade(merged, FILL_GREY)
    _write(merged, f"Total  ({result.scorable_weight:.1f} of {result.total_weight:g} weightage "
                   f"points auditable from the data supplied)", bold=True,
           align=WD_ALIGN_PARAGRAPH.RIGHT)
    for j, v in zip(range(5, 9), [f"{result.total_weight:g}",
                                  f"{result.index:.1f}%" if result.index else "—",
                                  f"{result.total_score:.1f}",
                                  f"Sales Process Index — {band['band']}"]):
        cell = row.cells[j]
        _shade(cell, FILL_GREY)
        _write(cell, v, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER if j < 8 else None)

    # --- Section B: pillar summary -------------------------------------------
    doc.add_page_break()
    heading(doc, "Section B — Process index summary by pillar")
    widths = [700, 5400, 1700, 1700, 1700, 4438]
    t = make_table(doc, widths, ["S.No", "Process pillar", "Weightage", "Score obtained",
                                 "Achievement %", "RAG status and key remark"])
    for i, p_ in enumerate(result.pillars, 1):
        row = t.add_row()
        worst = sorted([l for l in p_.lines if l.scored],
                       key=lambda l: l.achievement)[:1]
        remark = (f"Weakest line {worst[0].code} {worst[0].name} at "
                  f"{worst[0].achievement:.0f}%" if worst else "No line scorable")
        vals = [str(i), p_.title, f"{p_.weight:g} (scored {p_.scorable_weight:.1f})",
                f"{p_.score:.1f}",
                f"{p_.achievement:.1f}%" if p_.achievement is not None else "—",
                f"{p_.rag(cfg.rag)} — {remark}"]
        for j, v in enumerate(vals):
            cell = row.cells[j]
            _set_width(cell, widths[j])
            _write(cell, v, bold=(j in (0, 3)), size=8,
                   align=WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2, 3, 4) else None)
        fill = rag_fill(p_.rag(cfg.rag))
        if fill:
            _shade(row.cells[5], fill)
    row = t.add_row()
    merged = row.cells[0].merge(row.cells[1])
    _shade(merged, FILL_GREY)
    _write(merged, "SALES PROCESS INDEX", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    for j, v in zip(range(2, 6),
                    [f"{result.total_weight:g} (scored {result.scorable_weight:.1f})",
                     f"{result.total_score:.1f}",
                     f"{result.index:.1f}%" if result.index else "—",
                     f"{band['band']} — {band['response']}"]):
        _shade(row.cells[j], FILL_GREY)
        _write(row.cells[j], v, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER if j < 5 else None)

    para(doc, [(f"Sales Process Index = {result.total_score:.1f} ÷ {result.scorable_weight:.1f} "
                f"scorable weight = {result.index:.1f}% — {band['band']}. ", True),
               (f"{result.total_weight - result.scorable_weight:.1f} of "
                f"{result.total_weight:g} weightage points could not be scored because the "
                f"source data was not supplied; the index is calculated on the points that "
                f"were auditable and will change once the missing files are provided.", False)])

    # --- Section C: funnel ----------------------------------------------------
    heading(doc, "Section C — Funnel snapshot")
    f = ctx.get("funnel", [])
    if f:
        widths = [3400, 2600, 2200, 2600, 2200, 2638]
        t = make_table(doc, widths, ["Funnel stage", "Norm", "Count", "Loss / leakage",
                                     "Conversion", "Auditor remark"])
        for i, rowvals in enumerate(f):
            row = t.add_row()
            for j, v in enumerate(rowvals):
                cell = row.cells[j]
                _set_width(cell, widths[j])
                if i % 2:
                    _shade(cell, FILL_BAND)
                _write(cell, v, bold=(j == 0), size=8)

    # --- Section D: observations ---------------------------------------------
    doc.add_page_break()
    heading(doc, "Section D — Auditor observations")
    para(doc, [("Scope. ", True), (ctx.get("scope_text", ""), False)])
    if result.strengths:
        para(doc, [("Strengths observed", True)], space_after=2)
        for s in result.strengths:
            bullet(doc, s)
    if result.findings:
        para(doc, [("Critical findings", True)], space_after=2)
        for f_ in result.findings:
            bullet(doc, f_)
    if ctx.get("branch_notes"):
        para(doc, [("Branches requiring immediate attention", True)], space_after=2)
        for b in ctx["branch_notes"]:
            bullet(doc, b)

    # --- Section E: physical audit -------------------------------------------
    heading(doc, "Section E — Physical Audit Sheet")
    ph = ctx.get("physical_summary")
    if ph:
        widths = [4600, 1700, 1700, 1900, 5738]
        t = make_table(doc, widths, ["Section", "Items scored", "Failed",
                                     "Achievement", "Observation"])
        for i, rowvals in enumerate(ph):
            row = t.add_row()
            for j, v in enumerate(rowvals):
                cell = row.cells[j]
                _set_width(cell, widths[j])
                if i % 2:
                    _shade(cell, FILL_BAND)
                _write(cell, v, bold=(j == 0), size=8,
                       align=WD_ALIGN_PARAGRAPH.CENTER if j in (1, 2, 3) else None)
    else:
        para(doc, "Physical Audit Sheet not supplied, so pillar J is unscored. The "
                  "sheet covers showroom ambience, display and demo vehicle status, "
                  "uniform and grooming, the walk-in register and its reconciliation to DMS, "
                  "follow-up discipline, the sales consultant kit (current price list, "
                  "brochures, finance schemes, SHIELD and RSA tariff sheet), "
                  "customer amenities, mandatory displays, the delivery bay, safety and "
                  "statutory compliance, and performance visibility on the floor.")
        para(doc, [("To include it: ", True),
                   ("run  python run_audit.py --make-sheet  (or use the Download button on "
                    "the web app's upload page), complete one tab per branch during the "
                    "showroom walk, then save the file as Physical_Audit_Sheet.xlsx in the "
                    "input folder — or upload it in the same drop box as the DMS extracts. "
                    "Pillar J is then scored.", False)])

    # --- Section F: grading ---------------------------------------------------
    doc.add_page_break()
    heading(doc, "Section F — Grading and governance")
    widths = [2600, 3000, 3600, 6438]
    t = make_table(doc, widths, ["Process index", "Band", "Required response", "Governance action"])
    bands = cfg.grading
    for i, g in enumerate(bands):
        upper = bands[i - 1]["min"] - 0.1 if i else 100
        row = t.add_row()
        label = f"{g['min']} – {upper:g}" if i else f"{g['min']} – 100"
        for j, v in enumerate([label, g["band"], g["response"], g["action"]]):
            cell = row.cells[j]
            _set_width(cell, widths[j])
            if i % 2:
                _shade(cell, FILL_BAND)
            _write(cell, v, bold=(j == 1), size=8)
        if g["band"] == band["band"]:
            fill = rag_fill(g["band"].split(" ")[0])
            if fill:
                for c in row.cells:
                    _shade(c, fill)

    # --- Section G: corrective action plan -----------------------------------
    heading(doc, "Section G — Corrective action plan")
    para(doc, f"Every parameter below {cfg.rag['amber_min']}% achievement appears here with a "
              f"named owner and a dated commitment.", size=8)
    widths = [560, 1700, 3400, 3400, 1900, 1200, 3478]
    t = make_table(doc, widths, ["S.No", "Pillar", "Gap observed (with data)",
                                 "Corrective action agreed", "Owner", "Target date",
                                 "Input source and evidence"])
    for i, a in enumerate(result.actions, 1):
        row = t.add_row()
        for j, v in enumerate([str(i), a["pillar"], a["gap"], a["action"],
                               a["owner"], a["target"], a["source"]]):
            cell = row.cells[j]
            _set_width(cell, widths[j])
            if i % 2:
                _shade(cell, FILL_BAND)
            _write(cell, v, bold=(j == 0), size=8,
                   align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else None)

    # --- Section H: sign-off --------------------------------------------------
    heading(doc, "Section H — Sign-off")
    widths = [4000, 4200, 3600, 3838]
    t = make_table(doc, widths, ["Role", "Name", "Signature", "Date"])
    for i, role in enumerate(cfg.signoff):
        row = t.add_row()
        for j, v in enumerate([role, "", "", ""]):
            cell = row.cells[j]
            _set_width(cell, widths[j])
            if i % 2:
                _shade(cell, FILL_BAND)
            _write(cell, v, bold=(j == 0))
    para(doc, "Next audit due on: ______________          "
              "Re-audit required (index below 60): Yes / No", space_after=0)

    doc.save(path)
