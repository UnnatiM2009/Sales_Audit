"""Physical Audit Sheet — the observation-based half of the audit.

The DMS extracts cannot tell you whether the display cars are clean, whether
consultants carry a current price list, or whether the walk-in register is
being maintained. Those have to be walked and seen. This module:

  1. writes a blank branch-wise Physical Audit Sheet (`--make-sheet`)
  2. reads the completed sheet back and scores it as pillar J

Scoring per item: Yes = 2, Partial = 1, No = 0, NA = excluded from both
numerator and denominator so an irrelevant item does not penalise a branch.
"""
from __future__ import annotations

from dataclasses import dataclass

import pandas as pd

from .scoring import Line, Pillar, pct

RESPONSES = {"YES": 2.0, "PARTIAL": 1.0, "NO": 0.0}
NA_TOKENS = {"NA", "N/A", "NOT APPLICABLE", "-"}


@dataclass(frozen=True)
class CheckItem:
    section: str
    code: str
    item: str
    what_to_look_for: str
    weight: float = 1.0
    critical: bool = False   # a 'No' here is reported as a critical finding


# -----------------------------------------------------------------------------
# The audit sheet. Sections map to the pillar J sub-scores in the report.
# -----------------------------------------------------------------------------
CHECKLIST: tuple[CheckItem, ...] = (
    # --- 1. Showroom exterior and ambience -----------------------------------
    CheckItem("1. Exterior and ambience", "P1.1", "Glow sign and façade branding",
              "Signage lit, undamaged, current OEM corporate identity, no faded panels", 1),
    CheckItem("1. Exterior and ambience", "P1.2", "Approach, frontage and parking",
              "Frontage clean, customer parking marked and free of staff or stock vehicles", 1),
    CheckItem("1. Exterior and ambience", "P1.3", "Entrance and accessibility",
              "Entrance clear, ramp available, doors clean, no clutter at the threshold", 1),
    CheckItem("1. Exterior and ambience", "P1.4", "Showroom floor cleanliness",
              "Floor, glass, ceiling and lighting clean; all lights working; temperature comfortable", 1),
    CheckItem("1. Exterior and ambience", "P1.5", "Music, odour and general ambience",
              "No workshop noise or odour in the showroom; ambient music at appropriate volume", 1),

    # --- 2. Display vehicles --------------------------------------------------
    CheckItem("2. Display vehicles", "P2.1", "Display vehicle count against norm",
              "Number of display cars on the floor matches the OEM norm for this outlet", 2, True),
    CheckItem("2. Display vehicles", "P2.2", "Model and variant coverage",
              "Every live model represented; no discontinued variant on display", 2),
    CheckItem("2. Display vehicles", "P2.3", "Display vehicle cleanliness",
              "Exterior polished, interior vacuumed, glass clean, tyres dressed, no dust on dashboard", 2, True),
    CheckItem("2. Display vehicles", "P2.4", "Price and specification card on every car",
              "Current price card displayed on each vehicle, matching today's price list", 2, True),
    CheckItem("2. Display vehicles", "P2.5", "Display vehicle condition",
              "No scratches or dents, doors and boot open freely, seats unmarked, protective film removed", 1),
    CheckItem("2. Display vehicles", "P2.6", "Accessory and VAP display",
              "Accessories fitted on at least one display car; SHIELD and RSA offers displayed", 2, True),
    CheckItem("2. Display vehicles", "P2.7", "Battery, fuel and readiness",
              "Display cars start on demand; battery charged; infotainment demonstrable", 1),

    # --- 3. Demo and test drive vehicles -------------------------------------
    CheckItem("3. Demo and test drive fleet", "P3.1", "Demo vehicle availability per model",
              "At least one running demo per live model; physically present at the branch on audit day", 2, True),
    CheckItem("3. Demo and test drive fleet", "P3.2", "Statutory documents valid",
              "Insurance, PUC, fitness and registration valid and carried in the vehicle", 2, True),
    CheckItem("3. Demo and test drive fleet", "P3.3", "Demo vehicle cleanliness and condition",
              "Clean inside and out, fuel above half tank, no warning lights, tyres roadworthy", 2),
    CheckItem("3. Demo and test drive fleet", "P3.4", "Test drive route and consent",
              "Defined TD route displayed; customer licence verified and consent form signed before every drive", 1),
    CheckItem("3. Demo and test drive fleet", "P3.5", "Test drive register maintained",
              "Register filled for every drive with time out, time in, odometer and customer signature", 1),

    # --- 4. Uniform and grooming ----------------------------------------------
    CheckItem("4. Uniform and grooming", "P4.1", "Sales consultants in prescribed uniform",
              "All consultants on floor in current uniform, clean and well-fitted", 2, True),
    CheckItem("4. Uniform and grooming", "P4.2", "Name badge and identity card",
              "Every customer-facing employee wearing a legible name badge", 1),
    CheckItem("4. Uniform and grooming", "P4.3", "Grooming standard",
              "Hair, footwear and general presentation to dealership standard", 1),
    CheckItem("4. Uniform and grooming", "P4.4", "Support staff turnout",
              "Security, reception, driver and housekeeping in their respective uniforms", 1),

    # --- 5. Walk-in register and lead capture ---------------------------------
    CheckItem("5. Walk-in register", "P5.1", "Walk-in register maintained at entrance",
              "Physical or tablet register present and in active use by security or reception", 2, True),
    CheckItem("5. Walk-in register", "P5.2", "Register fields complete",
              "Name, mobile, model of interest, consultant assigned, in-time and out-time all filled", 2, True),
    CheckItem("5. Walk-in register", "P5.3", "Register reconciled to DMS same day",
              "Count in register for the last 3 days matches enquiries punched in DMS for those days", 2, True),
    CheckItem("5. Walk-in register", "P5.4", "Consultant allocation is rostered",
              "A visible allocation roster exists; walk-ins are not picked up ad hoc", 1),
    CheckItem("5. Walk-in register", "P5.5", "Non-converting walk-ins recorded",
              "Customers who leave without an enquiry are still logged with a reason", 1, True),

    # --- 6. Follow-up discipline ----------------------------------------------
    CheckItem("6. Follow-up discipline", "P6.1", "Follow-up calendar in use",
              "Each consultant can show today's follow-up list from the DMS or app", 2, True),
    CheckItem("6. Follow-up discipline", "P6.2", "Call remarks are substantive",
              "Sample 5 records per consultant: remarks describe the conversation, not 'called' or 'no response'", 2),
    CheckItem("6. Follow-up discipline", "P6.3", "Morning meeting held and recorded",
              "Minutes or attendance for the last 6 working days available", 1),
    CheckItem("6. Follow-up discipline", "P6.4", "Lost enquiry review",
              "Weekly lost-enquiry review with the branch manager, evidenced", 1),
    CheckItem("6. Follow-up discipline", "P6.5", "Aged enquiry ownership",
              "Enquiries older than 30 days have a named owner and a revival or closure plan", 1),

    # --- 7. Sales consultant kit ----------------------------------------------
    CheckItem("7. Sales consultant kit", "P7.1", "Current price list carried",
              "Sample 5 consultants: each has today's price list, matching the displayed cards", 2, True),
    CheckItem("7. Sales consultant kit", "P7.2", "Product brochures in stock",
              "Brochures for every live model available and not outdated", 1),
    CheckItem("7. Sales consultant kit", "P7.3", "Finance scheme sheet",
              "Current finance schemes, interest rates and EMI chart available", 1),
    CheckItem("7. Sales consultant kit", "P7.4", "SHIELD and RSA tariff sheet",
              "Tariff and plan comparison for SHIELD (extended warranty) and RSA carried by every consultant", 2, True),
    CheckItem("7. Sales consultant kit", "P7.5", "Accessory catalogue and price list",
              "Fitted-accessory catalogue with prices available at the desk", 1),
    CheckItem("7. Sales consultant kit", "P7.6", "Exchange evaluation format",
              "Blank evaluation forms and the evaluator's contact available on the floor", 1),
    CheckItem("7. Sales consultant kit", "P7.7", "Quotation format and calculator",
              "Standard quotation format in use; on-road price computed on the system, not by hand", 1),
    CheckItem("7. Sales consultant kit", "P7.8", "Product knowledge spot check",
              "Ask 2 consultants 3 product questions each; both answer correctly without help", 2),

    # --- 8. Customer amenities -------------------------------------------------
    CheckItem("8. Customer amenities", "P8.1", "Seating and lounge",
              "Adequate clean seating for waiting customers; upholstery undamaged", 1),
    CheckItem("8. Customer amenities", "P8.2", "Drinking water and refreshments",
              "Water available and offered; refreshment service functioning", 1),
    CheckItem("8. Customer amenities", "P8.3", "Washroom cleanliness",
              "Customer washroom clean, stocked, and checked on a signed roster", 2, True),
    CheckItem("8. Customer amenities", "P8.4", "Kids corner and Wi-Fi",
              "Kids area safe and clean; guest Wi-Fi available and working", 1),

    # --- 9. Mandatory displays and communication ------------------------------
    CheckItem("9. Mandatory displays", "P9.1", "Price list displayed publicly",
              "Current on-road price list on the customer notice board", 2, True),
    CheckItem("9. Mandatory displays", "P9.2", "SHIELD and RSA offer board",
              "Products, plans and prices displayed where the customer can read them unaided", 2, True),
    CheckItem("9. Mandatory displays", "P9.3", "Finance scheme and EMI board",
              "Tie-up financiers, rates and EMI illustrations displayed", 1),
    CheckItem("9. Mandatory displays", "P9.4", "Grievance officer and complaint box",
              "Name, number and escalation matrix displayed; complaint box present and opened on a schedule", 2, True),
    CheckItem("9. Mandatory displays", "P9.5", "OEM statutory and CI displays",
              "All OEM-mandated posters and statutory notices current and undamaged", 1),

    # --- 10. Delivery area -----------------------------------------------------
    CheckItem("10. Delivery experience", "P10.1", "Dedicated delivery bay",
              "Separate, clean, well-lit delivery bay not used for parking or storage", 2, True),
    CheckItem("10. Delivery experience", "P10.2", "Ceremonial delivery setup",
              "Decoration, lighting, photo backdrop and ribbon ready before the customer arrives", 1),
    CheckItem("10. Delivery experience", "P10.3", "Delivery checklist and HOTO station",
              "Pre-delivery checklist signed; HOTO form, washing checksheet and document folder ready", 2, True),
    CheckItem("10. Delivery experience", "P10.4", "SHIELD and RSA certificates in the folder",
              "Sample 3 recent deliveries: SHIELD and RSA certificates physically present", 2, True),
    CheckItem("10. Delivery experience", "P10.5", "Feature explanation at handover",
              "Consultant demonstrates key features and pairs the customer's phone before handover", 1),

    # --- 11. Systems, compliance and safety ------------------------------------
    CheckItem("11. Systems and safety", "P11.1", "DMS terminals and tablets working",
              "All sales terminals and tablets functional; no consultant working on paper", 1),
    CheckItem("11. Systems and safety", "P11.2", "EV charger operational",
              "Showroom charger working and accessible where EV or LMM models are sold", 1),
    CheckItem("11. Systems and safety", "P11.3", "Fire extinguishers and first aid",
              "Extinguishers in date and accessible; first aid box stocked", 2, True),
    CheckItem("11. Systems and safety", "P11.4", "CCTV functioning",
              "Cameras working with retention as per policy", 1),
    CheckItem("11. Systems and safety", "P11.5", "Statutory licences displayed",
              "Trade certificate, GST registration and shop licence current and displayed", 1),

    # --- 12. Performance visibility --------------------------------------------
    CheckItem("12. Performance visibility", "P12.1", "Target and achievement board",
              "Branch and consultant-wise target versus achievement updated to yesterday", 1),
    CheckItem("12. Performance visibility", "P12.2", "Funnel board",
              "Enquiry, test drive, booking and retail counts displayed and current", 1),
    CheckItem("12. Performance visibility", "P12.3", "SHIELD and RSA penetration board",
              "Penetration displayed by consultant so the gap is visible on the floor", 1, True),
    CheckItem("12. Performance visibility", "P12.4", "Previous audit actions displayed",
              "Open corrective actions from the last audit visible with owners and dates", 1),
)

SECTIONS = tuple(dict.fromkeys(c.section for c in CHECKLIST))


# -----------------------------------------------------------------------------
# Template
# -----------------------------------------------------------------------------
def write_template(path, branches: list[str],
                   answers: dict[str, dict[str, dict]] | None = None,
                   meta: dict[str, dict] | None = None) -> None:
    """Write a branch-wise Physical Audit Sheet.

    With `answers` it is pre-filled — `{branch: {code: {"r": "YES",
    "obs": ..., "owner": ..., "due": ...}}}` — which is how an audit captured
    on a phone becomes an ordinary sheet the scorer already reads. Without it
    the sheet is blank, for auditors who prefer Excel.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    notes = [
        ("PHYSICAL AUDIT SHEET — how to fill this", 14, True),
        ("", 10, False),
        ("One sheet per branch. Walk the showroom and complete every row.", 10, False),
        ("Response column accepts: Yes / Partial / No / NA", 10, False),
        ("   Yes     = fully compliant, evidence seen", 10, False),
        ("   Partial = present but incomplete, outdated or inconsistent", 10, False),
        ("   No      = absent or non-compliant", 10, False),
        ("   NA      = genuinely not applicable at this outlet (excluded from scoring)", 10, False),
        ("", 10, False),
        ("Record what you actually saw in the Observation column. 'Not maintained'", 10, False),
        ("is weaker evidence than 'register last filled 19-Aug, 41 walk-ins", 10, False),
        ("recorded against 118 enquiries punched'.", 10, False),
        ("", 10, False),
        ("Rows marked CRITICAL raise a corrective action automatically when", 10, False),
        ("answered No, even if the rest of that section scores well.", 10, False),
        ("", 10, False),
        ("WHERE TO PUT THIS FILE WHEN YOU ARE DONE", 11, True),
        ("", 10, False),
        ("Web app: upload it together with the DMS extracts, in the same drop", 10, False),
        ("box. Keep the filename Physical_Audit_Sheet.xlsx.", 10, False),
        ("", 10, False),
        ("Command line: save it into the input folder alongside Enquiry.xlsx,", 10, False),
        ("Booking.xlsx and the rest, then re-run. Pillar J will then be scored.", 10, False),
    ]
    for i, (text, size, bold) in enumerate(notes, start=1):
        c = ws.cell(i, 1, text)
        c.font = Font(name="Arial", size=size, bold=bold)
    ws.column_dimensions["A"].width = 95

    navy, band = "1F3864", "F2F5FA"
    for branch in branches:
        title = str(branch)[:28] or "Branch"
        s = wb.create_sheet(title)
        headers = ["Section", "Code", "Check item", "What to look for", "Weight",
                   "Critical", "Response", "Observation / evidence", "Owner", "Target date"]
        s.cell(1, 1, f"Physical Audit Sheet — {branch}").font = Font(
            name="Arial", size=13, bold=True, color=navy)
        m = (meta or {}).get(branch, {})
        if m.get("auditor") or m.get("audit_date"):
            byline = (f"Auditor: {m.get('auditor') or '______________'}     "
                      f"Date: {m.get('audit_date') or '____________'}     "
                      f"Captured on mobile     Branch manager: ______________")
        else:
            byline = ("Auditor: ______________     Date: ____________     "
                      "Branch manager: ______________")
        s.cell(2, 1, byline).font = Font(name="Arial", size=9, italic=True)
        for j, h in enumerate(headers, 1):
            c = s.cell(4, j, h)
            c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=navy)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        s.row_dimensions[4].height = 30

        filled = (answers or {}).get(branch, {})
        for i, item in enumerate(CHECKLIST):
            r = 5 + i
            got = filled.get(item.code) or {}
            resp = str(got.get("r", "") or "").title()      # YES -> Yes, NA -> Na
            if resp.upper() == "NA":
                resp = "NA"
            vals = [item.section, item.code, item.item, item.what_to_look_for,
                    item.weight, "CRITICAL" if item.critical else "", resp,
                    got.get("obs", ""), got.get("owner", ""), got.get("due", "")]
            for j, v in enumerate(vals, 1):
                c = s.cell(r, j, v)
                c.font = Font(name="Arial", size=10,
                              bold=(j == 6 and item.critical))
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2:
                    c.fill = PatternFill("solid", fgColor=band)

        dv = DataValidation(type="list", formula1='"Yes,Partial,No,NA"', allow_blank=True)
        dv.error = "Enter Yes, Partial, No or NA"
        s.add_data_validation(dv)
        dv.add(f"G5:G{4 + len(CHECKLIST)}")

        for j, w in enumerate([24, 8, 34, 56, 8, 10, 12, 46, 20, 14], 1):
            s.column_dimensions[get_column_letter(j)].width = w
        s.freeze_panes = s.cell(5, 1)

    wb.save(path)


# -----------------------------------------------------------------------------
# Scoring
# -----------------------------------------------------------------------------
def _normalise_response(v) -> str | None:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    s = str(v).strip().upper()
    if not s:
        return None
    if s in NA_TOKENS:
        return "NA"
    if s in RESPONSES:
        return s
    if s in {"Y", "TRUE", "1", "OK", "COMPLIANT"}:
        return "YES"
    if s in {"N", "FALSE", "0", "NOT OK", "NON-COMPLIANT"}:
        return "NO"
    if s.startswith("PART"):
        return "PARTIAL"
    return None


def score_physical(df_by_branch: dict[str, pd.DataFrame], weight: float,
                   cfg_rag: dict[str, int], source: str) -> tuple[Pillar, pd.DataFrame]:
    """Score the completed Physical Audit Sheet into pillar J plus a detail frame."""
    lookup = {c.code: c for c in CHECKLIST}
    rows: list[dict] = []

    for branch, df in df_by_branch.items():
        if df is None or df.empty:
            continue
        code_col = "Code" if "Code" in df.columns else df.columns[1]
        resp_col = "Response" if "Response" in df.columns else None
        obs_col = "Observation / evidence" if "Observation / evidence" in df.columns else None
        if resp_col is None:
            continue
        for _, r in df.iterrows():
            item = lookup.get(str(r.get(code_col, "")).strip())
            if item is None:
                continue
            resp = _normalise_response(r.get(resp_col))
            rows.append({
                "branch": branch,
                "section": item.section,
                "code": item.code,
                "item": item.item,
                "what": item.what_to_look_for,
                "weight": item.weight,
                "critical": item.critical,
                "response": resp or "Not answered",
                "points": RESPONSES.get(resp, 0.0) if resp and resp != "NA" else None,
                "max_points": item.weight * 2 if resp and resp != "NA" else None,
                "observation": str(r.get(obs_col, "") or "") if obs_col else "",
            })

    detail = pd.DataFrame(rows)
    pillar = Pillar(key="J_physical", title="J. Physical Audit Sheet", weight=weight)

    if detail.empty:
        for section in SECTIONS:
            items = [c for c in CHECKLIST if c.section == section]
            w = round(weight * sum(c.weight for c in items) /
                      sum(c.weight for c in CHECKLIST), 2)
            pillar.lines.append(Line(
                code=f"J{SECTIONS.index(section) + 1}",
                pillar="J_physical", name=section.split(". ", 1)[-1],
                checkpoint="; ".join(c.item for c in items[:3]) + (" …" if len(items) > 3 else ""),
                norm_text="All items compliant", weight=w,
                unscored_reason="Physical Audit Sheet not supplied "
                                "(run with --make-sheet to generate a blank one)",
                source="Not supplied"))
        return pillar, detail

    # Weight each item by its own weight; a 'Yes' scores 2, 'Partial' 1, 'No' 0.
    detail["earned"] = detail["weight"] * detail["points"]
    detail["possible"] = detail["max_points"]

    total_item_weight = sum(c.weight for c in CHECKLIST)
    for i, section in enumerate(SECTIONS, start=1):
        items = [c for c in CHECKLIST if c.section == section]
        sec_weight = round(weight * sum(c.weight for c in items) / total_item_weight, 2)
        sub = detail[detail["section"] == section]
        answered = sub[sub["possible"].notna()]
        if answered.empty:
            pillar.lines.append(Line(
                code=f"J{i}", pillar="J_physical", name=section.split(". ", 1)[-1],
                checkpoint="; ".join(c.item for c in items[:3]),
                norm_text="All items compliant", weight=sec_weight,
                unscored_reason="no responses recorded for this section",
                source=source))
            continue

        ach = pct(answered["earned"].sum(), answered["possible"].sum())
        fails = answered[answered["points"] == 0]
        crit = fails[fails["critical"]]
        branches_failing = sorted(set(fails["branch"]))
        remark_bits = []
        if len(crit):
            remark_bits.append(f"{len(crit)} critical item(s) failed")
        if branches_failing:
            remark_bits.append("branches: " + ", ".join(branches_failing[:4])
                               + (" …" if len(branches_failing) > 4 else ""))
        na = sub[sub["response"] == "NA"]
        actual = (f"{len(answered)} items scored across {answered['branch'].nunique()} branch(es); "
                  f"{len(fails)} failed, {len(answered[answered['points'] == 1])} partial"
                  + (f", {len(na)} marked NA" if len(na) else ""))

        pillar.lines.append(Line(
            code=f"J{i}", pillar="J_physical", name=section.split(". ", 1)[-1],
            checkpoint="; ".join(c.item for c in items[:3]) + (" …" if len(items) > 3 else ""),
            norm_text="All items compliant", weight=sec_weight,
            actual_text=actual, achievement=ach, source=source,
            remark="; ".join(remark_bits) if remark_bits else "all items compliant"))

    return pillar, detail


def is_sheet_file(name) -> bool:
    """Does this filename look like a Physical Audit Sheet?

    Deliberately loose, because the files arrive from a phone via WhatsApp or
    a browser download and pick up decorations on the way:
    `Physical_Audit_Sheet_YAVATMAL.xlsx`, `Physical Audit Sheet (1).xlsx`,
    `physical-audit-sheet-12-09-2026.xlsx` all mean the same thing. Being
    strict here is how a completed walk ends up silently unscored.
    """
    from pathlib import Path as _P
    stem = _P(str(name)).stem.lower()
    flat = "".join(ch if ch.isalnum() else " " for ch in stem)
    flat = " ".join(flat.split())
    return flat.startswith("physical audit")


def find_sheets(input_dir) -> list:
    """Every Physical Audit Sheet in a folder, sorted for a stable merge."""
    from pathlib import Path as _P
    folder = _P(input_dir)
    if not folder.is_dir():
        return []
    return sorted(p for p in folder.iterdir()
                  if p.is_file() and p.suffix.lower() in (".xlsx", ".xlsm", ".xls")
                  and is_sheet_file(p.name))


def sheet_answers(df: pd.DataFrame) -> dict[str, dict]:
    """One branch tab back into {code: {r, obs, owner, due}}."""
    out: dict[str, dict] = {}
    codes = {c.code for c in CHECKLIST}
    for _, row in df.iterrows():
        code = str(row.get("Code", "")).strip()
        if code not in codes:
            continue
        resp = _normalise_response(row.get("Response"))
        entry = {}
        if resp:
            entry["r"] = resp
        for key, col in (("obs", "Observation / evidence"), ("owner", "Owner"),
                         ("due", "Target date")):
            val = row.get(col)
            if val is not None and not (isinstance(val, float) and pd.isna(val)):
                text = str(val).strip()
                if text:
                    entry[key] = text
        if entry:
            out[code] = entry
    return out


def merge_sheets(paths: list, out_path) -> list[str]:
    """Combine several sheets into one workbook, one tab per branch.

    Auditing branch by branch on a phone produces one exported sheet per
    branch. Rather than making someone paste tabs together in Excel, they are
    merged here. Where two files carry the same branch the later one wins,
    which is what re-walking a branch to correct it should mean.
    """
    answers: dict[str, dict[str, dict]] = {}
    for path in paths:
        try:
            for branch, df in read_physical(path).items():
                answers.setdefault(str(branch), {}).update(sheet_answers(df))
        except Exception:  # noqa: BLE001 — one unreadable file must not sink the rest
            continue
    branches = [b for b in answers if answers[b]]
    if not branches:
        return []
    write_template(out_path, branches, answers=answers)
    return branches


def read_physical(path) -> dict[str, pd.DataFrame]:
    """Read a completed Physical Audit Sheet into {branch: dataframe}."""
    xl = pd.ExcelFile(path)
    out: dict[str, pd.DataFrame] = {}
    for sheet in xl.sheet_names:
        if sheet.lower().startswith("instruction"):
            continue
        # pandas treats "NA" and "N/A" as missing by default, which would turn
        # a deliberate 'not applicable' into 'not answered'. Both are excluded
        # from the score, but only one of them is a decision the auditor made,
        # and the report says so. Keep the literal text; only a blank is blank.
        df = pd.read_excel(xl, sheet, header=3, keep_default_na=False,
                           na_values=[""])
        df.columns = [str(c).strip() for c in df.columns]
        if "Code" in df.columns:
            out[sheet] = df
    return out


def critical_failures(detail: pd.DataFrame) -> pd.DataFrame:
    """Critical items answered No — these become corrective actions."""
    if detail.empty:
        return detail
    return detail[(detail["critical"]) & (detail["points"] == 0)].copy()
