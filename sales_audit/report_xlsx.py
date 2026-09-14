"""Findings workbook: the auditor's working paper and exception lists."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import INPUT_FILES, MISSING_DATA_NOTES, Config
from .loaders import AuditData
from .scoring import AuditResult

NAVY, BAND, GREY = "1F3864", "F2F5FA", "E8E8E8"
RED, AMBER, GREEN = "FFC7CE", "FFEB9C", "C6EFCE"
FONT = "Arial"
_thin = Side(style="thin", color="B8C1CC")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


class WorkbookWriter:
    def __init__(self) -> None:
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    def sheet(self, name: str, headers: list[str], rows: list[list],
              widths: list[int], title: str = "", note: str = "",
              rag_col: int | None = None) -> None:
        ws = self.wb.create_sheet(name[:31])
        r = 1
        if title:
            ws.cell(1, 1, title).font = Font(name=FONT, size=13, bold=True, color=NAVY)
            r = 2
        if note:
            c = ws.cell(r, 1, note)
            c.font = Font(name=FONT, size=9, italic=True, color="444444")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 30
            r += 1
        hr = r + 1
        for j, h in enumerate(headers, 1):
            c = ws.cell(hr, j, h)
            c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            c.border = BORDER
        ws.row_dimensions[hr].height = 32
        for i, row in enumerate(rows):
            for j, v in enumerate(row, 1):
                c = ws.cell(hr + 1 + i, j, v)
                c.font = Font(name=FONT, size=10)
                c.border = BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i % 2:
                    c.fill = PatternFill("solid", fgColor=BAND)
            if rag_col is not None and rag_col <= len(row):
                val = str(row[rag_col - 1])
                fill = (RED if val.startswith("Red") else AMBER if val.startswith("Amber")
                        else GREEN if val.startswith("Green") else None)
                if fill:
                    cc = ws.cell(hr + 1 + i, rag_col)
                    cc.fill = PatternFill("solid", fgColor=fill)
                    cc.font = Font(name=FONT, size=10, bold=True)
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(hr + 1, 1)

    def frame(self, name: str, df: pd.DataFrame, title: str = "", note: str = "",
              rag_col_name: str | None = None, max_width: int = 50) -> None:
        """Write a dataframe with sensible column widths."""
        if df is None or df.empty:
            self.sheet(name, ["Note"], [["No records met the criteria for this exception list."]],
                       [80], title=title, note=note)
            return
        headers = [str(c) for c in df.columns]
        rows = df.fillna("").astype(object).values.tolist()
        widths = []
        for c in df.columns:
            longest = max([len(str(c))] + [len(str(v)) for v in df[c].head(200)])
            widths.append(min(max(longest + 2, 10), max_width))
        rag_col = headers.index(rag_col_name) + 1 if rag_col_name in headers else None
        self.sheet(name, headers, rows, widths, title=title, note=note, rag_col=rag_col)

    def save(self, path: str | Path) -> None:
        self.wb.save(path)


def write_findings(path: str | Path, result: AuditResult, data: AuditData,
                   cfg: Config, physical_detail: pd.DataFrame | None = None) -> None:
    w = WorkbookWriter()
    ex = result.exceptions
    ctx = result.context

    # --- 0. Source register ---------------------------------------------------
    rows = []
    for spec in INPUT_FILES:
        ds = data.ds(spec.ref)
        if ds:
            rows.append([spec.ref, ds.filename, ds.sheet, f"{ds.rows:,}",
                         len(ds.df.columns), spec.label, spec.scores, "Supplied"])
        else:
            rows.append([spec.ref, spec.pattern, "—", "—", "—", spec.label,
                         spec.scores, "Not supplied"])
    for key, label in MISSING_DATA_NOTES.items():
        if cfg.target(key) or cfg.norm(key):
            continue
        rows.append(["—", label, "—", "—", "—", "Reference data",
                     "Lines left unscored without it", "Not supplied"])
    w.sheet("0. Source register",
            ["Ref", "Input source file", "Sheet", "Records", "Columns",
             "What it is", "What it scores", "Status"],
            rows, [7, 46, 22, 12, 10, 34, 34, 14],
            title="Input source files used for this audit",
            note=(f"Audit period {ctx.get('period', '—')}. "
                  f"Sales Process Index {ctx.get('index_text', '—')}. "
                  "Every figure in this workbook traces to a file listed here."))

    # --- 1. Scorecard ---------------------------------------------------------
    rows = []
    for p in result.pillars:
        rows.append([f"{p.title}  —  weightage {p.weight:g}", "", "", "", "", "", "", ""])
        for l in p.lines:
            rows.append([l.code, l.name, l.checkpoint, l.norm_text, l.actual_text,
                         f"{l.weight:.1f}", l.display_achievement(), l.display_score(),
                         l.rag_remark(cfg.rag), l.source])
    w.sheet("1. Scorecard",
            ["S.No", "Audit parameter", "Checkpoint", "Norm", "Actual observed",
             "Wt", "Ach %", "Score", "RAG and remark", "Input source"],
            [r if len(r) == 10 else r + [""] * (10 - len(r)) for r in rows],
            [7, 26, 34, 26, 52, 7, 8, 8, 46, 44],
            title="Audit scorecard — every parameter",
            note=(f"Index {ctx.get('index_text', '—')}. "
                  f"{result.scorable_weight:.1f} of {result.total_weight:g} weightage points "
                  "were auditable from the data supplied."),
            rag_col=9)

    # --- 2. Pillar summary ----------------------------------------------------
    rows = [[p.title, f"{p.weight:g}", f"{p.scorable_weight:.1f}", f"{p.score:.1f}",
             f"{p.achievement:.1f}%" if p.achievement is not None else "—",
             p.rag(cfg.rag)] for p in result.pillars]
    rows.append(["SALES PROCESS INDEX", f"{result.total_weight:g}",
                 f"{result.scorable_weight:.1f}", f"{result.total_score:.1f}",
                 f"{result.index:.1f}%" if result.index is not None else "—",
                 result.band(cfg.grading)["band"]])
    w.sheet("2. Pillar summary",
            ["Pillar", "Weightage", "Scorable", "Score", "Achievement", "RAG"],
            rows, [46, 12, 11, 10, 14, 26],
            title="Index summary by pillar", rag_col=6)

    # --- exception tabs -------------------------------------------------------
    w.frame("3. Funnel by branch", ex.get("branch_funnel"),
            title="Full funnel by branch",
            note="Network conversion is the TOTAL row. Compare each branch against it.")
    w.frame("4. VAP RSA EW audit", ex.get("vap_fields"),
            title="RSA, extended warranty and VAP — field population on invoices",
            note=("A zero fill rate on SHIELD or RSA means either the product is not sold "
                  "or it is sold outside the DMS. Establish which before setting targets."),
            rag_col_name="RAG")
    w.frame("5. Retail exceptions", ex.get("retail_exceptions"),
            title="Invoices breaching a delivery control",
            note="Hand this list to the branch managers named in the corrective action plan.")
    w.frame("6. SC exceptions", ex.get("sc_exceptions"),
            title="Consultants below the conversion or productivity norm",
            note="Filtered to consultants holding 20 or more enquiries in the period.")
    w.frame("7. Aged enquiries", ex.get("aged_enquiries"),
            title="Live enquiries ageing beyond the working threshold",
            note="Each must be re-contacted and revived, or closed with a reason code.")
    w.frame("8. Complaints", ex.get("complaints"),
            title="Customer concerns across all funnel stages",
            note=("Rows marked 'Other outlet' fall outside the branch list in norms.yaml "
                  "and are excluded from the complaints-per-1000 rates."))
    w.frame("9. Duplicate phones", ex.get("duplicate_phones"),
            title="Same mobile number against multiple enquiries",
            note="Numbers worked by more than one consultant are an ownership risk.",
            rag_col_name="Risk")

    # --- 10. Physical audit ---------------------------------------------------
    if physical_detail is not None and not physical_detail.empty:
        det = physical_detail.rename(columns={
            "branch": "Branch", "section": "Section", "code": "Code", "item": "Check item",
            "what": "What to look for", "response": "Response",
            "observation": "Observation / evidence"})
        det["RAG"] = det["points"].map({2.0: "Green", 1.0: "Amber", 0.0: "Red"}).fillna("—")
        cols = ["Branch", "Section", "Code", "Check item", "What to look for",
                "Response", "RAG", "Observation / evidence"]
        w.frame("10. Physical Audit Sheet", det[cols],
                title="Physical Audit Sheet — item by item",
                note="Ambience, display and demo vehicles, uniforms, walk-in register, "
                     "consultant kit, mandatory displays, delivery bay and safety.",
                rag_col_name="RAG")
        crit = det[(det["critical"]) & (det["points"] == 0)] if "critical" in det else pd.DataFrame()
        if not crit.empty:
            w.frame("11. Physical critical fails", crit[cols],
                    title="Critical Physical Audit Sheet items answered No",
                    note="Each of these raises a corrective action automatically.",
                    rag_col_name="RAG")
    else:
        w.sheet("10. Physical Audit Sheet", ["Note"],
                [["Physical Audit Sheet not supplied, so pillar J is unscored."],
                 ["1. Generate a blank sheet:  python run_audit.py --make-sheet   "
                  "(or use the Download button on the web app's upload page)."],
                 ["2. Complete one tab per branch during the showroom walk."],
                 ["3. Save it as Physical_Audit_Sheet.xlsx in the input folder — or, on "
                  "the web app, upload it in the same drop box as the DMS extracts."],
                 ["Pillar J will then be scored and this tab will list every item."]],
                [118], title="Physical Audit Sheet")

    # --- corrective action plan ----------------------------------------------
    rows = [[str(i + 1), a["pillar"], a["gap"], a["action"], a["owner"],
             a["target"], "", a["source"]] for i, a in enumerate(result.actions)]
    w.sheet("12. Corrective action plan",
            ["S.No", "Pillar", "Gap observed (with data)", "Corrective action",
             "Owner", "Target date", "Status / closure evidence", "Input source"],
            rows, [7, 24, 64, 64, 28, 14, 26, 44],
            title="Corrective action plan",
            note="Every gap below is traceable to a tab in this workbook. "
                 "Owner and target date to be confirmed at sign-off.")

    w.save(path)
